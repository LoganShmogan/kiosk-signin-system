"""
Tests for export_util.py — exercises row-building logic and file generation.
"""
import os
import export_util


def _sign(member_id, first, last, direction, timestamp, signature=None):
    return {
        "member_id": member_id,
        "first_name": first,
        "last_name": last,
        "direction": direction,
        "timestamp": timestamp,
        "signature": signature,
    }


def _staff_sign(member_id, first, last, direction, timestamp, role="", signature=None):
    return {
        "member_id": member_id,
        "first_name": first,
        "last_name": last,
        "role": role,
        "direction": direction,
        "timestamp": timestamp,
        "signature": signature,
    }


# ── render_signature ──────────────────────────────────────────────────────────

def test_render_signature_returns_none_for_empty_bytes():
    assert export_util.render_signature(None) is None
    assert export_util.render_signature(b"") is None


def test_render_signature_returns_none_for_garbage_bytes():
    assert export_util.render_signature(b"not a python literal") is None


def test_render_signature_returns_png_bytes_for_valid_strokes():
    png = export_util.render_signature(b"[(0, 0, 10, 10), (10, 10, 20, 0)]")
    assert png is not None
    assert png[:8] == b"\x89PNG\r\n\x1a\n"


def test_render_signature_returns_none_for_empty_strokes_list():
    assert export_util.render_signature(b"[]") is None


def test_render_signature_returns_none_when_pil_unavailable(monkeypatch):
    monkeypatch.setattr(export_util, "PIL_AVAILABLE", False)
    assert export_util.render_signature(b"[(0, 0, 10, 10)]") is None


def test_pil_unavailable_falls_back_when_import_fails(monkeypatch):
    """Covers the `except ImportError: PIL_AVAILABLE = False` guard around
    the top-of-module `from PIL import Image, ImageDraw`."""
    import sys
    import importlib

    original_pil = sys.modules.get("PIL")
    monkeypatch.setitem(sys.modules, "PIL", None)
    try:
        importlib.reload(export_util)
        assert export_util.PIL_AVAILABLE is False
    finally:
        if original_pil is not None:
            sys.modules["PIL"] = original_pil
        else:
            sys.modules.pop("PIL", None)
        importlib.reload(export_util)
        assert export_util.PIL_AVAILABLE is True


# ── generate_export ────────────────────────────────────────────────────────────
# generate_export shares its layout with generate_staff_export (#78): one
# worksheet per member, each holding that member's sign-in/out sessions as
# Date | Sign In | Sign Out | Duration rows.

def test_generate_export_creates_xlsx(tmp_path):
    signs = [
        _sign(1, "Alice", "Brown", "IN",  "2026-01-01 08:30:00"),
        _sign(1, "Alice", "Brown", "OUT", "2026-01-01 14:00:00"),
    ]
    out = str(tmp_path / "export.xlsx")
    result = export_util.generate_export(signs, "2026-01-01", "2026-01-01", out)
    assert result == out
    assert os.path.exists(out)
    assert os.path.getsize(out) > 0


def test_generate_export_to_bytesio():
    import io
    signs = [
        _sign(1, "Alice", "Brown", "IN",  "2026-01-01 08:30:00"),
        _sign(1, "Alice", "Brown", "OUT", "2026-01-01 14:00:00"),
    ]
    buf = io.BytesIO()
    export_util.generate_export(signs, "2026-01-01", "2026-01-01", buf)
    buf.seek(0)
    data = buf.read()
    assert len(data) > 0
    assert data[:2] == b"PK"


def test_generate_export_one_sheet_per_member(tmp_path):
    """Regression test for #78: member export uses the same one-sheet-per-
    person layout as the staff export, instead of one sheet per calendar day."""
    import openpyxl

    signs = [
        _sign(1, "Alice", "Brown",  "IN",  "2026-01-01 08:00:00"),
        _sign(1, "Alice", "Brown",  "OUT", "2026-01-01 15:00:00"),
        _sign(2, "Bob",   "Carter", "IN",  "2026-01-01 08:05:00"),
    ]
    out = str(tmp_path / "export_multi.xlsx")
    export_util.generate_export(signs, "2026-01-01", "2026-01-01", out)

    wb = openpyxl.load_workbook(out)
    assert set(wb.sheetnames) == {"Alice Brown", "Bob Carter"}


def test_generate_export_multi_day_history_stays_on_one_sheet(tmp_path):
    """A member's sign-ins across several days stay on their one sheet, as
    separate session rows, rather than being split across per-day sheets."""
    import openpyxl

    signs = [
        _sign(1, "Alice", "Brown", "IN",  "2026-01-01 08:30:00"),
        _sign(1, "Alice", "Brown", "OUT", "2026-01-01 15:00:00"),
        _sign(1, "Alice", "Brown", "IN",  "2026-01-02 08:45:00"),
        _sign(1, "Alice", "Brown", "OUT", "2026-01-02 15:10:00"),
        _sign(1, "Alice", "Brown", "IN",  "2026-01-05 08:15:00"),
        _sign(1, "Alice", "Brown", "OUT", "2026-01-05 15:20:00"),
    ]
    out = str(tmp_path / "export_range.xlsx")
    export_util.generate_export(signs, "2026-01-01", "2026-01-05", out)

    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames == ["Alice Brown"]
    values = [c.value for row in wb["Alice Brown"].iter_rows() for c in row if c.value]
    assert "08:30" in values and "08:45" in values and "08:15" in values


def test_generate_export_empty_signs_list(tmp_path):
    out = str(tmp_path / "export_empty.xlsx")
    export_util.generate_export([], "2026-01-01", "2026-01-05", out)
    assert os.path.exists(out)
    assert os.path.getsize(out) > 0


def test_generate_export_embeds_signature_image(tmp_path):
    """A sign event with real stroke data must get an embedded PNG on the
    signature row underneath its session's data row."""
    import openpyxl

    signs = [
        _sign(1, "Alice", "Brown", "IN",  "2026-01-01 08:30:00",
              signature=b"[(0, 0, 10, 10), (10, 10, 20, 0)]"),
        _sign(1, "Alice", "Brown", "OUT", "2026-01-01 15:00:00"),
    ]
    out = str(tmp_path / "export_with_sig.xlsx")
    export_util.generate_export(signs, "2026-01-01", "2026-01-01", out)

    wb = openpyxl.load_workbook(out)
    ws = wb["Alice Brown"]
    assert len(ws._images) >= 1


# ── _build_staff_sessions ──────────────────────────────────────────────────────

def test_build_staff_sessions_empty():
    assert export_util._build_staff_sessions([]) == []


def test_build_staff_sessions_single_session():
    signs = [
        _staff_sign(1, "Alice", "Smith", "IN",  "2026-06-08 08:30:00", role="Coordinator"),
        _staff_sign(1, "Alice", "Smith", "OUT", "2026-06-08 11:45:00", role="Coordinator"),
    ]
    sections = export_util._build_staff_sessions(signs)
    assert len(sections) == 1
    key, sessions = sections[0]
    assert key[1] == "Alice"
    assert key[2] == "Smith"
    assert key[3] == "Coordinator"
    assert len(sessions) == 1
    sess = sessions[0]
    assert sess["date"] == "2026-06-08"
    assert sess["in"]["timestamp"][11:16]  == "08:30"
    assert sess["out"]["timestamp"][11:16] == "11:45"


def test_build_staff_sessions_multiple_sessions_same_day():
    # Staff can sign in/out any number of times, any time of day.
    signs = [
        _staff_sign(1, "Alice", "Smith", "IN",  "2026-06-08 08:00:00"),
        _staff_sign(1, "Alice", "Smith", "OUT", "2026-06-08 11:30:00"),
        _staff_sign(1, "Alice", "Smith", "IN",  "2026-06-08 13:00:00"),
        _staff_sign(1, "Alice", "Smith", "OUT", "2026-06-08 17:30:00"),
        _staff_sign(1, "Alice", "Smith", "IN",  "2026-06-08 20:00:00"),
        _staff_sign(1, "Alice", "Smith", "OUT", "2026-06-08 21:15:00"),
    ]
    _, sessions = export_util._build_staff_sessions(signs)[0]
    assert len(sessions) == 3
    assert [s["in"]["timestamp"][11:16] for s in sessions] == ["08:00", "13:00", "20:00"]
    assert [s["out"]["timestamp"][11:16] for s in sessions] == ["11:30", "17:30", "21:15"]


def test_build_staff_sessions_multiple_days():
    signs = [
        _staff_sign(1, "Alice", "Smith", "IN",  "2026-06-08 08:00:00"),
        _staff_sign(1, "Alice", "Smith", "OUT", "2026-06-08 16:00:00"),
        _staff_sign(1, "Alice", "Smith", "IN",  "2026-06-09 09:00:00"),
        _staff_sign(1, "Alice", "Smith", "OUT", "2026-06-09 16:00:00"),
    ]
    _, sessions = export_util._build_staff_sessions(signs)[0]
    assert [s["date"] for s in sessions] == ["2026-06-08", "2026-06-09"]


def test_build_staff_sessions_sorted_by_last_name():
    signs = [
        _staff_sign(1, "Zoe",   "Zebra",  "IN", "2026-06-08 08:00:00"),
        _staff_sign(2, "Alice", "Aardvark","IN", "2026-06-08 08:05:00"),
    ]
    sections = export_util._build_staff_sessions(signs)
    assert sections[0][0][2] == "Aardvark"
    assert sections[1][0][2] == "Zebra"


def test_build_staff_sessions_still_signed_in():
    # A dangling IN with no matching OUT yet is still included, with out=None.
    signs = [_staff_sign(1, "Alice", "Smith", "IN", "2026-06-08 08:00:00")]
    _, sessions = export_util._build_staff_sessions(signs)[0]
    assert len(sessions) == 1
    assert sessions[0]["in"] is not None
    assert sessions[0]["out"] is None


def test_build_staff_sessions_stray_out_with_no_in():
    # A stray OUT with nothing open still produces a session (in=None).
    signs = [_staff_sign(1, "Alice", "Smith", "OUT", "2026-06-08 11:00:00")]
    _, sessions = export_util._build_staff_sessions(signs)[0]
    assert len(sessions) == 1
    assert sessions[0]["in"] is None
    assert sessions[0]["out"] is not None


def test_build_staff_sessions_double_in_closes_previous_with_no_out():
    signs = [
        _staff_sign(1, "Alice", "Smith", "IN", "2026-06-08 08:00:00"),
        _staff_sign(1, "Alice", "Smith", "IN", "2026-06-08 09:00:00"),
        _staff_sign(1, "Alice", "Smith", "OUT", "2026-06-08 17:00:00"),
    ]
    _, sessions = export_util._build_staff_sessions(signs)[0]
    assert len(sessions) == 2
    assert sessions[0]["in"]["timestamp"][11:16] == "08:00"
    assert sessions[0]["out"] is None
    assert sessions[1]["in"]["timestamp"][11:16] == "09:00"
    assert sessions[1]["out"]["timestamp"][11:16] == "17:00"


def test_build_staff_sessions_multiple_staff():
    signs = [
        _staff_sign(1, "Alice", "Smith", "IN", "2026-06-08 08:00:00"),
        _staff_sign(2, "Bob",   "Jones", "IN", "2026-06-08 08:05:00"),
    ]
    sections = export_util._build_staff_sessions(signs)
    assert len(sections) == 2


def test_build_staff_sessions_empty_role():
    signs = [_staff_sign(1, "Alice", "Smith", "IN", "2026-06-08 08:00:00", role="")]
    key, _ = export_util._build_staff_sessions(signs)[0]
    assert key[3] == ""


# ── _day_month / _year_label ────────────────────────────────────────────────────

def test_day_month_formats_without_year():
    assert export_util._day_month("2026-06-08") == "08/06"


def test_day_month_malformed_date_returns_original_string():
    assert export_util._day_month("not-a-date") == "not-a-date"


def test_year_label_single_year():
    sessions = [{"date": "2026-06-08"}, {"date": "2026-11-02"}]
    assert export_util._year_label(sessions) == "2026"


def test_year_label_spans_years():
    sessions = [{"date": "2026-12-20"}, {"date": "2027-01-05"}]
    assert export_util._year_label(sessions) == "2026–2027"


def test_year_label_empty():
    assert export_util._year_label([]) == ""


# ── _unique_sheet_title ──────────────────────────────────────────────────────

def test_unique_sheet_title_basic():
    used = set()
    assert export_util._unique_sheet_title("Alice Brown", used) == "Alice Brown"


def test_unique_sheet_title_dedupes_collisions():
    used = set()
    first = export_util._unique_sheet_title("Alice Brown", used)
    second = export_util._unique_sheet_title("Alice Brown", used)
    assert first == "Alice Brown"
    assert second == "Alice Brown (2)"


def test_unique_sheet_title_strips_invalid_excel_characters():
    used = set()
    title = export_util._unique_sheet_title("A/B:C*D?E[F]G\\H", used)
    assert not set(title) & set("[]:*?/\\")


def test_unique_sheet_title_truncates_to_31_chars():
    used = set()
    title = export_util._unique_sheet_title("A" * 50, used)
    assert len(title) <= 31


# ── _duration_mins / _fmt_duration ────────────────────────────────────────────

def _evt(ts):
    """Minimal sign event with just a timestamp."""
    return {"timestamp": ts, "signature": None}


def test_duration_mins_normal():
    assert export_util._duration_mins(_evt("2026-06-09 08:00:00"),
                                      _evt("2026-06-09 11:30:00")) == 210


def test_duration_mins_one_hour():
    assert export_util._duration_mins(_evt("2026-06-09 13:00:00"),
                                      _evt("2026-06-09 14:00:00")) == 60


def test_duration_mins_missing_start():
    assert export_util._duration_mins(None, _evt("2026-06-09 11:00:00")) is None


def test_duration_mins_missing_finish():
    assert export_util._duration_mins(_evt("2026-06-09 08:00:00"), None) is None


def test_duration_mins_both_missing():
    assert export_util._duration_mins(None, None) is None


def test_duration_mins_malformed_timestamp_returns_none():
    assert export_util._duration_mins(_evt("not-a-timestamp"),
                                      _evt("2026-06-09 11:00:00")) is None


def test_fmt_duration_hours_and_minutes():
    assert export_util._fmt_duration(90)  == "1h 30m"
    assert export_util._fmt_duration(210) == "3h 30m"


def test_fmt_duration_exact_hours():
    assert export_util._fmt_duration(120) == "2h 00m"


def test_fmt_duration_less_than_one_hour():
    assert export_util._fmt_duration(45) == "0h 45m"


def test_fmt_duration_none():
    assert export_util._fmt_duration(None) == ""


# ── generate_staff_export ──────────────────────────────────────────────────────

def test_generate_staff_export_creates_xlsx(tmp_path):
    signs = [
        _staff_sign(1, "Alice", "Smith", "IN",  "2026-06-09 08:30:00", role="Manager"),
        _staff_sign(1, "Alice", "Smith", "OUT", "2026-06-09 17:00:00", role="Manager"),
    ]
    out = str(tmp_path / "staff_export.xlsx")
    result = export_util.generate_staff_export(signs, "2026-06-09", "2026-06-13", out)
    assert result == out
    assert os.path.exists(out)
    assert os.path.getsize(out) > 0


def test_generate_staff_export_to_bytesio():
    import io
    signs = [
        _staff_sign(1, "Alice", "Smith", "IN",  "2026-06-09 08:30:00", role="Manager"),
        _staff_sign(1, "Alice", "Smith", "OUT", "2026-06-09 17:00:00", role="Manager"),
    ]
    buf = io.BytesIO()
    export_util.generate_staff_export(signs, "2026-06-09", "2026-06-13", buf)
    buf.seek(0)
    data = buf.read()
    assert len(data) > 0
    assert data[:2] == b"PK"


def test_generate_staff_export_multiple_staff(tmp_path):
    signs = [
        _staff_sign(1, "Alice", "Smith", "IN",  "2026-06-09 08:00:00", role="Manager"),
        _staff_sign(2, "Bob",   "Jones", "IN",  "2026-06-09 09:00:00", role=""),
        _staff_sign(2, "Bob",   "Jones", "OUT", "2026-06-09 17:00:00", role=""),
    ]
    out = str(tmp_path / "staff_multi.xlsx")
    export_util.generate_staff_export(signs, "2026-06-09", "2026-06-13", out)
    assert os.path.exists(out)
    assert os.path.getsize(out) > 0


def test_generate_staff_export_one_sheet_per_staff_member(tmp_path):
    """Regression test for #76: downloading all staff gives each staff member
    their own worksheet, instead of one shared sheet with page breaks."""
    import openpyxl

    signs = [
        _staff_sign(1, "Alice", "Smith", "IN",  "2026-06-09 08:00:00", role="Manager"),
        _staff_sign(1, "Alice", "Smith", "OUT", "2026-06-09 17:00:00", role="Manager"),
        _staff_sign(2, "Bob",   "Jones", "IN",  "2026-06-09 09:00:00", role=""),
        _staff_sign(2, "Bob",   "Jones", "OUT", "2026-06-09 17:00:00", role=""),
    ]
    out = str(tmp_path / "staff_sheets.xlsx")
    export_util.generate_staff_export(signs, "2026-06-09", "2026-06-13", out)

    wb = openpyxl.load_workbook(out)
    assert set(wb.sheetnames) == {"Alice Smith", "Bob Jones"}


def test_generate_staff_export_no_signs_for_some_days(tmp_path):
    signs = [
        _staff_sign(1, "Alice", "Smith", "IN", "2026-06-08 08:00:00"),
        _staff_sign(1, "Alice", "Smith", "IN", "2026-06-10 08:00:00"),
    ]
    out = str(tmp_path / "staff_sparse.xlsx")
    export_util.generate_staff_export(signs, "2026-06-08", "2026-06-12", out)
    assert os.path.exists(out)


def test_generate_staff_export_empty_signs_list(tmp_path):
    out = str(tmp_path / "staff_empty.xlsx")
    export_util.generate_staff_export([], "2026-06-08", "2026-06-12", out)
    assert os.path.exists(out)
    assert os.path.getsize(out) > 0


def test_generate_staff_export_still_signed_in(tmp_path):
    """A staff member with no sign-out yet must not crash the export."""
    import openpyxl

    signs = [_staff_sign(1, "Alice", "Smith", "IN", "2026-06-09 08:00:00", role="Manager")]
    out = str(tmp_path / "staff_open.xlsx")
    export_util.generate_staff_export(signs, "2026-06-09", "2026-06-09", out)

    wb = openpyxl.load_workbook(out)
    ws = wb.active
    values = [c.value for row in ws.iter_rows() for c in row if c.value]
    assert any("Still signed in" in str(v) for v in values)


def test_generate_staff_export_embeds_signature_image(tmp_path):
    """A session with a captured signature gets an embedded PNG on the
    extra signature row underneath its data row."""
    import openpyxl

    signs = [
        _staff_sign(1, "Alice", "Smith", "IN",  "2026-06-09 08:00:00",
                    role="Manager", signature=b"[(0, 0, 10, 10), (10, 10, 20, 0)]"),
        _staff_sign(1, "Alice", "Smith", "OUT", "2026-06-09 16:00:00", role="Manager"),
    ]
    out = str(tmp_path / "staff_with_sig.xlsx")
    export_util.generate_staff_export(signs, "2026-06-09", "2026-06-09", out)

    wb = openpyxl.load_workbook(out)
    ws = wb.active
    assert len(ws._images) >= 1


def test_generate_staff_export_totals_correct(tmp_path):
    """The per-person total at the bottom sums every session's duration."""
    import openpyxl

    signs = [
        # 08:00 – 11:30 = 210 min = 3h 30m
        _staff_sign(1, "Alice", "Smith", "IN",  "2026-06-09 08:00:00", role="Manager"),
        _staff_sign(1, "Alice", "Smith", "OUT", "2026-06-09 11:30:00", role="Manager"),
        # 13:00 – 17:00 = 240 min = 4h 00m
        _staff_sign(1, "Alice", "Smith", "IN",  "2026-06-09 13:00:00", role="Manager"),
        _staff_sign(1, "Alice", "Smith", "OUT", "2026-06-09 17:00:00", role="Manager"),
    ]
    out = str(tmp_path / "staff_totals.xlsx")
    export_util.generate_staff_export(signs, "2026-06-09", "2026-06-09", out)

    wb = openpyxl.load_workbook(out)
    ws = wb.active
    values = [c.value for row in ws.iter_rows() for c in row if c.value]
    assert any("Total time in the centre: 7h 30m" == v for v in values)


def test_generate_staff_export_no_sessions_shows_placeholder(tmp_path, monkeypatch):
    """Defensive branch: if a (member, sessions) pair with an empty sessions
    list ever reaches the renderer, it shows a 'no records' row instead of
    crashing. `_build_staff_sessions` always produces at least one session per
    member for real input, so this is exercised via monkeypatch."""
    import openpyxl

    monkeypatch.setattr(
        export_util, "_build_staff_sessions",
        lambda signs: [((1, "Alice", "Smith", "Manager"), [])],
    )
    out = str(tmp_path / "staff_no_sessions.xlsx")
    export_util.generate_staff_export([], "2026-06-09", "2026-06-09", out)

    wb = openpyxl.load_workbook(out)
    ws = wb.active
    values = [c.value for row in ws.iter_rows() for c in row if c.value]
    assert any("No sign-in records for this period." in str(v) for v in values)


def test_generate_staff_export_year_header(tmp_path):
    import openpyxl

    signs = [
        _staff_sign(1, "Alice", "Smith", "IN",  "2026-06-09 08:00:00", role="Manager"),
        _staff_sign(1, "Alice", "Smith", "OUT", "2026-06-09 17:00:00", role="Manager"),
    ]
    out = str(tmp_path / "staff_year.xlsx")
    export_util.generate_staff_export(signs, "2026-06-09", "2026-06-09", out)

    wb = openpyxl.load_workbook(out)
    ws = wb.active
    values = [c.value for row in ws.iter_rows() for c in row if c.value]
    assert any("Alice Smith" in str(v) and "2026" in str(v) for v in values)
