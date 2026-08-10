import ast
import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

try:
    from PIL import Image, ImageDraw
    PIL_AVAILABLE = True
except ImportError:
    PIL_AVAILABLE = False


def render_signature(sig_bytes: bytes, width=160, height=60) -> bytes | None:
    """Convert stroke data (stored as a serialised list) into a PNG byte string."""
    if not sig_bytes or not PIL_AVAILABLE:
        return None
    try:
        strokes = ast.literal_eval(sig_bytes.decode())
    except Exception:
        return None
    if not strokes:
        return None

    xs = [c for s in strokes for c in (s[0], s[2])]
    ys = [c for s in strokes for c in (s[1], s[3])]
    min_x, max_x = min(xs), max(xs)
    min_y, max_y = min(ys), max(ys)
    range_x = max(max_x - min_x, 1)
    range_y = max(max_y - min_y, 1)

    margin  = 6
    scale   = min((width - margin * 2) / range_x, (height - margin * 2) / range_y)

    def tx(x): return int((x - min_x) * scale) + margin
    def ty(y): return int((y - min_y) * scale) + margin

    img  = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(img)
    for (x1, y1, x2, y2) in strokes:
        draw.line([(tx(x1), ty(y1)), (tx(x2), ty(y2))], fill="black", width=2)

    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


DATA_FONT   = Font(name="Arial", size=11)
ALT_FILL    = PatternFill("solid", start_color="F0EDE8")
THIN        = Side(style="thin", color="D0CBE8")
BORDER      = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER      = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT        = Alignment(horizontal="left",   vertical="center", wrap_text=True)
ROW_HEIGHT  = 45


_INVALID_SHEET_CHARS = set('[]:*?/\\')


def _unique_sheet_title(label: str, used: set) -> str:
    """An Excel-safe worksheet title (<=31 chars, no reserved characters),
    de-duplicated against titles already used in this workbook."""
    base = "".join(c for c in label if c not in _INVALID_SHEET_CHARS)[:31] or "Sheet"
    title = base
    n = 2
    while title in used:
        suffix = f" ({n})"
        title = base[:31 - len(suffix)] + suffix
        n += 1
    used.add(title)
    return title


# ── Person export (shared by members and staff) ────────────────────────────────
# Sign in/out with no restriction on time of day or count — so each row is one
# sign-in/out session (not a fixed AM/PM slot), in the order it happened. Every
# person gets their own worksheet, so bulk downloads print/browse cleanly.

STAFF_HEADERS    = ["Date", "Sign In", "Sign Out", "Duration"]
STAFF_COL_WIDTHS = [10, 14, 18, 12]

STAFF_NAME_FILL   = PatternFill("solid", start_color="372976")
STAFF_NAME_FONT   = Font(name="Arial", bold=True, size=12, color="FFFFFF")
STAFF_SUBHDR_FILL = PatternFill("solid", start_color="4a3a8a")
STAFF_SUBHDR_FONT = Font(name="Arial", bold=True, size=10, color="FFFFFF")
TOTAL_FONT        = Font(name="Arial", bold=True, size=11)
STILL_IN_FONT     = Font(name="Arial", size=11, italic=True, color="BB2233")

SIG_ROW_HEIGHT = 34


def _duration_mins(start_event, finish_event):
    """Minutes between two sign events, or None if either is missing."""
    if not start_event or not finish_event:
        return None
    try:
        t1 = datetime.strptime(start_event["timestamp"], "%Y-%m-%d %H:%M:%S")
        t2 = datetime.strptime(finish_event["timestamp"], "%Y-%m-%d %H:%M:%S")
        mins = int((t2 - t1).total_seconds() / 60)
        return mins if mins >= 0 else None
    except Exception:
        return None


def _fmt_duration(mins) -> str:
    """Format minutes as 'Xh YYm', or '' if None."""
    if mins is None:
        return ""
    return f"{mins // 60}h {mins % 60:02d}m"


def _day_month(date_str: str) -> str:
    """'YYYY-MM-DD' -> 'DD/MM' (no year — the year is shown once, per person)."""
    try:
        return datetime.strptime(date_str, "%Y-%m-%d").strftime("%d/%m")
    except Exception:
        return date_str


def _year_label(sessions: list) -> str:
    """Single year ('2026') or a range ('2026–2027') covering all sessions."""
    years = sorted({s["date"][:4] for s in sessions if s.get("date")})
    if not years:
        return ""
    return years[0] if len(years) == 1 else f"{years[0]}–{years[-1]}"


def _build_staff_sessions(signs: list) -> list:
    """
    Groups raw sign events by person and pairs them into sessions.

    Returns a list of (member_key, sessions) tuples sorted by last name, where
    member_key = (member_id, first_name, last_name, role) and sessions is a
    chronologically ordered list of dicts:
        {"date": "YYYY-MM-DD", "in": event_or_None, "out": event_or_None}

    Sign-ins/outs aren't restricted to a time of day or a fixed count per day,
    so sessions are built by simple alternation: the first IN opens a session
    and the next OUT closes it. A stray IN while one is already open closes
    the previous session with no sign-out; a stray OUT with nothing open opens
    a session with no sign-in. A session still open at the end of the data
    (signed in, not yet signed out) is included with out=None.
    """
    by_member: dict = {}
    order: list = []
    for s in signs:
        key = (s["member_id"], s["first_name"], s["last_name"], s.get("role") or "")
        if key not in by_member:
            by_member[key] = []
            order.append(key)
        by_member[key].append(s)

    result = []
    for key in order:
        events = sorted(by_member[key], key=lambda e: e["timestamp"])
        sessions = []
        open_in = None
        for e in events:
            if e["direction"] == "IN":
                if open_in is not None:
                    sessions.append({"date": open_in["timestamp"][:10],
                                      "in": open_in, "out": None})
                open_in = e
            else:  # OUT
                if open_in is not None:
                    sessions.append({"date": open_in["timestamp"][:10],
                                      "in": open_in, "out": e})
                    open_in = None
                else:
                    sessions.append({"date": e["timestamp"][:10],
                                      "in": None, "out": e})
        if open_in is not None:
            sessions.append({"date": open_in["timestamp"][:10], "in": open_in, "out": None})
        result.append((key, sessions))

    result.sort(key=lambda x: (x[0][2].lower(), x[0][1].lower()))
    return result


def _write_person_sheet(ws, key, sessions, date_from, date_to, title_prefix) -> None:
    """Populate one worksheet with a single person's sign-in/out sessions.

    Portrait layout: one row per sign-in/out session (Date | Sign In | Sign Out |
    Duration), with a second row underneath for signature images when present.
    A name/role/year banner sits up top and a total time-in-centre line at
    the bottom.
    """
    _, fname, lname, role = key
    last_col = get_column_letter(len(STAFF_HEADERS))

    for col_idx, width in enumerate(STAFF_COL_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    # ── Title ──────────────────────────────────────────────────────────────────
    ws.merge_cells(f"A1:{last_col}1")
    ws["A1"]           = f"{title_prefix}  |  {date_from}  to  {date_to}"
    ws["A1"].font      = Font(name="Arial", bold=True, size=13, color="372976")
    ws["A1"].alignment = CENTER
    ws.row_dimensions[1].height = 28

    current_row = 2

    # ── Person name banner ─────────────────────────────────────────────────────
    name = f"{fname} {lname}".strip() or "Unknown"
    parts = [name]
    if role:
        parts.append(role)
    year_label = _year_label(sessions)
    if year_label:
        parts.append(year_label)
    banner_text = "  —  ".join(parts)
    ws.merge_cells(f"A{current_row}:{last_col}{current_row}")
    c           = ws[f"A{current_row}"]
    c.value     = banner_text
    c.font      = STAFF_NAME_FONT
    c.fill      = STAFF_NAME_FILL
    c.alignment = LEFT
    ws.row_dimensions[current_row].height = 26
    current_row += 1

    # ── Column headers ─────────────────────────────────────────────────────────
    for col_idx, hdr in enumerate(STAFF_HEADERS, start=1):
        cell           = ws.cell(row=current_row, column=col_idx, value=hdr)
        cell.font      = STAFF_SUBHDR_FONT
        cell.fill      = STAFF_SUBHDR_FILL
        cell.alignment = CENTER
        cell.border    = BORDER
    ws.row_dimensions[current_row].height = 20
    current_row += 1

    # ── One (or two) rows per session ───────────────────────────────────────────
    total_mins = 0
    if not sessions:
        ws.merge_cells(f"A{current_row}:{last_col}{current_row}")
        mc           = ws[f"A{current_row}"]
        mc.value     = "No sign-in records for this period."
        mc.font      = Font(name="Arial", size=10, italic=True, color="999999")
        mc.alignment = CENTER
        current_row += 1
    else:
        for s_idx, sess in enumerate(sessions):
            in_evt, out_evt = sess["in"], sess["out"]
            mins = _duration_mins(in_evt, out_evt)
            if mins is not None:
                total_mins += mins

            row_fill = ALT_FILL if s_idx % 2 == 0 else None
            ws.row_dimensions[current_row].height = ROW_HEIGHT

            in_text  = in_evt["timestamp"][11:16] if in_evt else ""
            if out_evt is not None:
                out_text, out_font = out_evt["timestamp"][11:16], DATA_FONT
            else:
                out_text, out_font = "Still signed in", STILL_IN_FONT

            values = [
                (_day_month(sess["date"]), DATA_FONT, LEFT),
                (in_text,                  DATA_FONT, CENTER),
                (out_text,                 out_font,  CENTER),
                (_fmt_duration(mins),      TOTAL_FONT, CENTER),
            ]
            for col_idx, (val, font, align) in enumerate(values, start=1):
                cell           = ws.cell(row=current_row, column=col_idx, value=val)
                cell.font      = font
                cell.alignment = align
                cell.border    = BORDER
                if row_fill:
                    cell.fill = row_fill
            current_row += 1

            in_sig  = in_evt["signature"]  if in_evt  else None
            out_sig = out_evt["signature"] if out_evt else None
            if in_sig or out_sig:
                ws.row_dimensions[current_row].height = SIG_ROW_HEIGHT
                for col_idx in range(1, len(STAFF_HEADERS) + 1):
                    cell        = ws.cell(row=current_row, column=col_idx, value="")
                    cell.border = BORDER
                    if row_fill:
                        cell.fill = row_fill
                for col_idx, sig_bytes in ((2, in_sig), (3, out_sig)):
                    png = render_signature(sig_bytes) if sig_bytes else None
                    if png:
                        from openpyxl.drawing.image import Image as XLImage
                        img_obj        = XLImage(io.BytesIO(png))
                        img_obj.width  = 100
                        img_obj.height = 32
                        ws.add_image(img_obj, f"{get_column_letter(col_idx)}{current_row}")
                current_row += 1

    # ── Total time in the centre ─────────────────────────────────────────────────
    ws.merge_cells(f"A{current_row}:{last_col}{current_row}")
    tc           = ws[f"A{current_row}"]
    tc.value     = f"Total time in the centre: {_fmt_duration(total_mins)}"
    tc.font      = TOTAL_FONT
    tc.fill      = ALT_FILL
    tc.alignment = LEFT
    ws.row_dimensions[current_row].height = 22

    from openpyxl.worksheet.page import PageMargins
    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize   = 9
    ws.page_setup.fitToPage   = True
    ws.page_setup.fitToWidth  = 1
    ws.page_setup.fitToHeight = 0
    ws.page_margins = PageMargins(left=0.4, right=0.4, top=0.6, bottom=0.6)


def _generate_person_sessions_export(signs: list, date_from: str, date_to: str,
                                     output_path, title_prefix: str) -> str:
    """Build a .xlsx with one worksheet per person, each showing their
    sign-in/out sessions. Shared by generate_export and generate_staff_export
    so member and staff exports look identical and every bulk download gives
    each person their own sheet."""
    sessions_by_person = _build_staff_sessions(signs)

    wb = Workbook()
    wb.remove(wb.active)
    used_titles = set()

    if not sessions_by_person:
        ws = wb.create_sheet(title=_unique_sheet_title("No Records", used_titles))
        last_col = get_column_letter(len(STAFF_HEADERS))
        for col_idx, width in enumerate(STAFF_COL_WIDTHS, start=1):
            ws.column_dimensions[get_column_letter(col_idx)].width = width
        ws.merge_cells(f"A1:{last_col}1")
        ws["A1"]           = f"{title_prefix}  |  {date_from}  to  {date_to}"
        ws["A1"].font      = Font(name="Arial", bold=True, size=13, color="372976")
        ws["A1"].alignment = CENTER
        ws.row_dimensions[1].height = 28
        ws.merge_cells(f"A2:{last_col}2")
        ws["A2"]           = "No sign-in records for this period."
        ws["A2"].font      = Font(name="Arial", size=10, italic=True, color="999999")
        ws["A2"].alignment = CENTER
    else:
        for key, sessions in sessions_by_person:
            _, fname, lname, _ = key
            ws = wb.create_sheet(title=_unique_sheet_title(f"{fname} {lname}".strip(), used_titles))
            _write_person_sheet(ws, key, sessions, date_from, date_to, title_prefix)

    wb.save(output_path)
    return output_path


def generate_export(signs: list, date_from: str, date_to: str, output_path) -> str:
    """Build the member sign-in/out .xlsx. output_path can be a file path or file-like object.

    Same format as generate_staff_export: one worksheet per member, each
    holding that person's sign-in/out sessions.
    """
    return _generate_person_sessions_export(signs, date_from, date_to, output_path,
                                            title_prefix="Kiosk Sign-In Export")


def generate_staff_export(signs: list, date_from: str, date_to: str, output_path) -> str:
    """Build the staff sign-in/out .xlsx. output_path can be a file path or file-like object.

    One worksheet per staff member, each holding that person's sign-in/out
    sessions.
    """
    return _generate_person_sessions_export(signs, date_from, date_to, output_path,
                                            title_prefix="Kiosk Staff Sign-In Export")
